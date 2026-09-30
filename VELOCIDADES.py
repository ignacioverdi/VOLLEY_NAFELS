# -*- coding: utf-8 -*-
"""
VELOCIDADES.py — la velocidad de los saques, de la pistola a la app

COMO SE USA
-----------
  1. Arrastra el .dvw sobre VELOCIDADES.bat.
     Te arma velocidades_<PARTIDO>.xlsx con una fila por saque, ya llena:
     numero, set, minuto del video, dorsal, jugador, tipo y valoracion.
     Vos solo escribis los km/h en la columna amarilla.

  2. Cuando terminaste de cargar, arrastra el mismo .dvw otra vez.
     Lee la planilla y deja velocidades_<PARTIDO>.json al lado.

  3. HACER_TODO agarra ese .json solo y le pega la velocidad a cada saque.
     El jugador la ve arriba a la derecha del video.

POR QUE LA PLANILLA SE GENERA Y NO SE ESCRIBE A MANO
----------------------------------------------------
Una planilla escrita a mano hay que mantenerla alineada con el .dvw: si el
scout cargo un saque de mas, o el orden no es el que se creia, todas las
velocidades se corren de lugar y quedan puestas al jugador equivocado. Paso una
vez y no se noto hasta compararlo con el video.

Generandola desde el .dvw eso no puede pasar: cada fila ya trae su jugador y su
minuto. Si una lectura no salio, se deja la celda vacia y las demas no se
mueven.
"""
import io
import json
import os
import re
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print()
    print('     Falta openpyxl (la libreria que lee y escribe Excel).')
    print('     Instalala una vez con:   pip install openpyxl')
    input('\n     Enter para cerrar. ')
    sys.exit(1)

try:
    from CORTAR_SAQUES import leer_dvw, pedir, mmss
except ImportError:
    print('     Tiene que estar CORTAR_SAQUES.py en la misma carpeta.')
    input('\n     Enter para cerrar. ')
    sys.exit(1)

TIPO = {'M': 'flotado', 'Q': 'potencia', 'T': 'salto flotado', 'H': 'alta',
        'O': 'otro', 'N': 'sin tipo'}
VAL = {'#': 'ace', '/': 'sin ataque', '+': 'buena', '!': 'neutra',
       '-': 'mala', '=': 'error'}

AMARILLO = PatternFill('solid', fgColor='FFF3B0')
GRIS = PatternFill('solid', fgColor='EFEFEF')
TITULO = Font(bold=True, size=13)
CAB = Font(bold=True, color='FFFFFF')
CABFILL = PatternFill('solid', fgColor='1F3B57')
FINA = Side(style='thin', color='BBBBBB')
BORDE = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)


def slug(dvw):
    b = os.path.splitext(os.path.basename(dvw))[0]
    b = re.sub(r'^[&\s]+', '', b)
    return re.sub(r'[^A-Za-z0-9]+', '_', b).strip('_').upper()[:40] or 'PARTIDO'


# ══════════════════════════════════════════════════════════════════════════
#   armar la planilla
# ══════════════════════════════════════════════════════════════════════════

def armar(saques, equipos, ruta, nombre):
    w = Workbook()
    s = w.active
    s.title = 'SAQUES'

    s['A1'] = 'VELOCIDAD DE SAQUE — %s' % nombre
    s['A1'].font = TITULO
    s['A2'] = ('%d saques. Escribi SOLO la columna amarilla (KM/H). '
               'Si la pistola no leyo, deja la celda vacia: NO borres la fila.'
               % len(saques))
    s['A3'] = ('Las demas columnas salen del scout y estan para que sepas a que '
               'saque le estas poniendo el numero.')
    for f in ('A2', 'A3'):
        s[f].font = Font(italic=True, color='555555')

    cab = ['#', 'SET', 'MIN', 'EQUIPO', 'DORSAL', 'JUGADOR', 'TIPO',
           'RESULTADO', 'KM/H']
    for j, c in enumerate(cab, 1):
        z = s.cell(row=5, column=j, value=c)
        z.font = CAB; z.fill = CABFILL
        z.alignment = Alignment(horizontal='center')
        z.border = BORDE

    for i, x in enumerate(saques):
        r = 6 + i
        vals = [x['n'], x['set'], mmss(x['seg']), x['equipo'], x['num'],
                (x['ape'] + (' ' + x['nombre'] if x['nombre'] else '')).strip(),
                TIPO.get(x['tipo'], x['tipo']), VAL.get(x['val'], x['val']), None]
        for j, v in enumerate(vals, 1):
            z = s.cell(row=r, column=j, value=v)
            z.border = BORDE
            if j == 9:
                z.fill = AMARILLO
                z.number_format = '0.0'
                z.alignment = Alignment(horizontal='center')
            else:
                z.fill = GRIS
                if j in (1, 2, 5):
                    z.alignment = Alignment(horizontal='center')

    for j, an in enumerate([6, 6, 9, 20, 8, 26, 14, 14, 10], 1):
        s.column_dimensions[get_column_letter(j)].width = an
    s.freeze_panes = 'A6'

    # ── la hoja que se calcula sola ──
    r2 = w.create_sheet('RESUMEN')
    r2['A1'] = 'PROMEDIOS POR JUGADOR'
    r2['A1'].font = TITULO
    r2['A2'] = 'Se calcula solo a medida que cargas las velocidades.'
    r2['A2'].font = Font(italic=True, color='555555')
    cab2 = ['JUGADOR', 'DORSAL', 'SAQUES', 'PROMEDIO', 'MAXIMO',
            'PROM. POTENCIA', 'PROM. FLOTADO', 'PROM. ACES']
    for j, c in enumerate(cab2, 1):
        z = r2.cell(row=4, column=j, value=c)
        z.font = CAB; z.fill = CABFILL
        z.alignment = Alignment(horizontal='center'); z.border = BORDE

    ult = 5 + len(saques)
    dor = 'SAQUES!$E$6:$E$%d' % ult
    kmh = 'SAQUES!$I$6:$I$%d' % ult
    tip = 'SAQUES!$G$6:$G$%d' % ult
    res = 'SAQUES!$H$6:$H$%d' % ult
    vistos = []
    for x in saques:
        k = (x['num'], (x['ape'] + (' ' + x['nombre'] if x['nombre'] else '')).strip())
        if k not in vistos:
            vistos.append(k)
    for i, (num, nom) in enumerate(sorted(vistos)):
        r = 5 + i
        r2.cell(row=r, column=1, value=nom).border = BORDE
        r2.cell(row=r, column=2, value=num).border = BORDE
        f = [
            '=COUNTIFS(%s,$B%d,%s,">0")' % (dor, r, kmh),
            '=IFERROR(AVERAGEIFS(%s,%s,$B%d,%s,">0"),"")' % (kmh, dor, r, kmh),
            '=IFERROR(SUMPRODUCT(MAX((%s=$B%d)*(%s>0)*%s)),"")' % (dor, r, kmh, kmh),
            '=IFERROR(AVERAGEIFS(%s,%s,$B%d,%s,"potencia",%s,">0"),"")' % (kmh, dor, r, tip, kmh),
            '=IFERROR(AVERAGEIFS(%s,%s,$B%d,%s,"flotado",%s,">0"),"")' % (kmh, dor, r, tip, kmh),
            '=IFERROR(AVERAGEIFS(%s,%s,$B%d,%s,"ace",%s,">0"),"")' % (kmh, dor, r, res, kmh),
        ]
        for j, ff in enumerate(f, 3):
            z = r2.cell(row=r, column=j, value=ff)
            z.border = BORDE
            z.alignment = Alignment(horizontal='center')
            if j > 3:
                z.number_format = '0.0'
    for j, an in enumerate([26, 8, 9, 11, 10, 16, 16, 12], 1):
        r2.column_dimensions[get_column_letter(j)].width = an
    r2.freeze_panes = 'A5'

    w.save(ruta)


# ══════════════════════════════════════════════════════════════════════════
#   leerla de vuelta
# ══════════════════════════════════════════════════════════════════════════

def leer(ruta, saques):
    w = load_workbook(ruta, data_only=True)
    s = w['SAQUES']
    por_n = {x['n']: x for x in saques}
    out = {}
    sin_leer = 0
    for r in s.iter_rows(min_row=6, values_only=True):
        if not r or r[0] is None:
            continue
        try:
            n = int(r[0])
        except (TypeError, ValueError):
            continue
        x = por_n.get(n)
        if not x:
            continue
        v = r[8] if len(r) > 8 else None
        if not isinstance(v, (int, float)) or not (10 < float(v) < 200):
            sin_leer += 1
            continue
        out[n] = dict(saque=n, set=x['set'], equipo=x['equipo'], num=x['num'],
                      ape=x['ape'], nombre=x['nombre'], tipo=x['tipo'],
                      val=x['val'], seg=x['seg'], kmh=round(float(v), 1),
                      fuente='pistola')
    return out, sin_leer


def main():
    print()
    print('  ' + '=' * 68)
    print('     VELOCIDADES DE SAQUE — de la pistola a la app')
    print('  ' + '=' * 68)

    args = [a for a in sys.argv[1:] if os.path.isfile(a)]
    dvw = args[0] if args else pedir('\n     Arrastra el .dvw y apreta Enter: ')
    if not dvw or not os.path.exists(dvw):
        print('     No encuentro ese archivo.')
        input('\n     Enter para cerrar. ')
        return

    _v, saques, equipos = leer_dvw(dvw)
    if not saques:
        print('     Ese .dvw no tiene saques con tiempo de video.')
        input('\n     Enter para cerrar. ')
        return

    carpeta = os.path.dirname(os.path.abspath(dvw))
    base = slug(dvw)
    xlsx = os.path.join(carpeta, 'velocidades_%s.xlsx' % base)
    jsn = os.path.join(carpeta, 'velocidades_%s.json' % base)
    nombre = os.path.splitext(os.path.basename(dvw))[0].lstrip('& ')

    print()
    print('     %s  vs  %s' % (equipos['*'], equipos['a']))
    print('     %d saques.' % len(saques))

    if not os.path.exists(xlsx):
        armar(saques, equipos, xlsx, nombre)
        print()
        print('     Te deje la planilla:')
        print('       %s' % os.path.basename(xlsx))
        print()
        print('     Abrila, escribi los km/h en la columna amarilla y guardala.')
        print('     Cuando termines, arrastra este mismo .dvw otra vez y yo')
        print('     me encargo del resto.')
        print()
        print('     (la hoja RESUMEN se calcula sola mientras cargas)')
        print('  ' + '=' * 68)
        input('\n     Enter para cerrar. ')
        return

    print()
    print('     Leyendo %s...' % os.path.basename(xlsx))
    try:
        cargados, vacias = leer(xlsx, saques)
    except Exception as e:
        print('     No pude leer la planilla: %s' % e)
        print('     Fijate que no la tengas abierta en Excel.')
        input('\n     Enter para cerrar. ')
        return

    if not cargados:
        print()
        print('     La planilla todavia no tiene ninguna velocidad cargada.')
        print('     Escribi los km/h en la columna amarilla, guardala, y volve.')
        input('\n     Enter para cerrar. ')
        return

    doc = dict(partido=base, origen='pistola de radar (medicion manual)',
               desfase_aplicado=0.0,
               medidos=len(cargados), de_saques=len(saques),
               saques={str(k): v for k, v in sorted(cargados.items())})
    with io.open(jsn, 'w', encoding='utf-8') as f:
        f.write(json.dumps(doc, ensure_ascii=False, indent=1))

    v = sorted(x['kmh'] for x in cargados.values())
    print()
    print('     %d velocidades cargadas de %d saques.' % (len(cargados), len(saques)))
    if vacias:
        print('     (%d filas sin velocidad: quedan sin numero, no molestan)' % vacias)
    print('     Mas rapido %.1f   ·   promedio %.1f   ·   mas lento %.1f km/h'
          % (v[-1], sum(v) / len(v), v[0]))
    porj = {}
    for x in cargados.values():
        porj.setdefault('%s #%s' % (x['ape'], x['num']), []).append(x['kmh'])
    print()
    for n in sorted(porj, key=lambda a: -sum(porj[a]) / len(porj[a])):
        z = porj[n]
        print('       %-22s %5.1f km/h de promedio  (max %.1f, %d saques)'
              % (n, sum(z) / len(z), max(z), len(z)))
    print()
    print('     Guardado en  %s' % os.path.basename(jsn))
    print()
    print('     Ahora corre HACER_TODO y la velocidad le queda pegada a cada')
    print('     saque. El jugador la ve arriba a la derecha del video.')
    print('  ' + '=' * 68)
    input('\n     Enter para cerrar. ')


if __name__ == '__main__':
    try:
        main()
    except KeyboardInterrupt:
        print('\n     Chau.')
